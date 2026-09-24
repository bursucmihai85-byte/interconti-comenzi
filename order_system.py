import sys
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import re
import json
import os
import urllib.parse
from datetime import datetime
from rapidfuzz import fuzz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

JSON_PATH = os.path.join(os.path.dirname(__file__), "nomenclator.json")

with open(JSON_PATH, "r", encoding="utf-8") as f:
    NOMENCLATOR = json.load(f)

BOGDAN_PHONE = "40733959906"

# Inițializare Gemini AI
GEMINI_CLIENT = None
try:
    from google import genai
    from google.genai import types
    env_file = os.path.join(os.path.dirname(__file__), ".env")
    gemini_key = os.environ.get("GEMINI_API_KEY", "")
    if not gemini_key and os.path.exists(env_file):
        for line in open(env_file, "r", encoding="utf-8"):
            if line.strip().startswith("GEMINI_API_KEY="):
                gemini_key = line.split("=", 1)[1].strip()
                break
    if gemini_key:
        GEMINI_CLIENT = genai.Client(api_key=gemini_key)
        print("-> [AI] Gemini Client initializat cu succes in order_system!")
except Exception as e:
    print(f"-> [AI] Initializare Gemini esuata: {e}")

PRODUCT_TYPES = [
    r'ramifica[tț]i[ie]', r'ramifica[tț]ii',
    r'[tț]eav[aă]', r'[tț]evi',
    r'niplu', r'nipluri',
    r'cot', r'coturi',
    r'teu', r'teuri',
    r'muf[aă]', r'mufe',
    r'robinet', r'robine[tț]i',
    r'sifon', r'[sș]ifoane', r'[sș]efu(\s*[\-]?\s*[iî]n)?',
    r'aerisitor', r'aerisitoare',
    r'filtru', r'filtre',
    r'colier', r'coliere',
    r'reduc[tț]ie', r'reduc[tț]ii',
    r'dop', r'dopuri',
    r'capac', r'capace',
    r'calorifer', r'calorifere',
    r'radiator', r'radiatoare',
    r'pomp[aă]', r'pompe',
    r'racord', r'racorduri',
    r'supap[aă]', r'supape',
    r'clapet', r'clapete',
    r'distribuitor', r'distribuitoare',
    r'olandez', r'olandezi',
    r'ventil', r'garnitur[aă]', r'garnituri'
]

PROD_REGEX = r'(?:\b(?:' + '|'.join(PRODUCT_TYPES) + r')\b)'
QTY_REGEX = r'(\b\d+\s*(?:buc(?:ati|ăți)?|m(?:etri)?|ml|set(?:uri)?|rol(?:e|a)?|cut(?:ii|ie)?)\b)'

def normalize_plurals(word: str) -> str:
    w = word.upper()
    repl = [
        (r'RAMIFICA[TȚ]II$', 'RAMIFICATIE'),
        (r'[TȚ]EVI$', 'TEAVA'),
        (r'NIPLURI$', 'NIPLU'),
        (r'COTURI$', 'COT'),
        (r'TEURI$', 'TEU'),
        (r'MUFE$', 'MUFA'),
        (r'ROBINE[TȚ]I$', 'ROBINET'),
        (r'[SȘ]IFOANE$', 'SIFON'),
        (r'AERISITOARE$', 'AERISITOR'),
        (r'FILTRE$', 'FILTRU'),
        (r'COLIERE$', 'COLIER'),
        (r'REDUC[TȚ]II$', 'REDUCTIE'),
        (r'DOPURI$', 'DOP'),
        (r'CAPACE$', 'CAPAC'),
        (r'CALORIFERE$', 'CALORIFER'),
        (r'RADIATOARE$', 'RADIATOR'),
        (r'POMPE$', 'POMPA'),
        (r'RACORDURI$', 'RACORD'),
        (r'SUPAPE$', 'SUPAPA'),
        (r'OLANDEZI$', 'OLANDEZ'),
        (r'GARNITURI$', 'GARNITURA')
    ]
    for p, r in repl:
        w = re.sub(p, r, w)
    return w

def parse_multi_item_dictation(text: str) -> list[str]:
    t = text.strip()
    t = re.sub(r'\b(cam\s+at[aâ]t|at[aâ]t|gata|mul[tț]umesc)\b.*$', '', t, flags=re.IGNORECASE).strip()
    
    # 1. Separare pe virgule și conjuncții ("și", "apoi", "plus")
    t = re.sub(r'\s*,\s*', '\n', t)
    t = re.sub(r'\s+\b(și|si|apoi|plus)\b\s+', '\n', t, flags=re.IGNORECASE)
    
    initial_segments = [s.strip() for s in t.split('\n') if s.strip()]
    
    stage2_segments = []
    for s in initial_segments:
        sub_s = re.sub(
            r'(' + PROD_REGEX + r'.*?\s+' + QTY_REGEX + r')\s+(?=(?:\d+\s*(?:buc|m|ml|set|rol|cut)?\s+)?' + PROD_REGEX + ')',
            r'\1\n',
            s,
            flags=re.IGNORECASE
        )
        for part in sub_s.split('\n'):
            if part.strip():
                stage2_segments.append(part.strip())

    final_items = []
    for s in stage2_segments:
        matches = list(re.finditer(PROD_REGEX, s, flags=re.IGNORECASE))
        if len(matches) <= 1:
            final_items.append(s)
        else:
            split_indices = []
            for i in range(1, len(matches)):
                idx = matches[i].start()
                prefix = s[:idx]
                qty_match = re.search(r'(\b\d+\s*(?:buc(?:ati|ăți)?|m(?:etri)?|ml|set(?:uri)?|rol(?:e|a)?|cut(?:ii|ie)?)?\s*)$', prefix)
                if qty_match:
                    idx -= len(qty_match.group(1))
                split_indices.append(idx)
            
            last = 0
            for idx in split_indices:
                sub = s[last:idx].strip()
                if sub:
                    final_items.append(sub)
                last = idx
            sub = s[last:].strip()
            if sub:
                final_items.append(sub)
                
    results = [r for r in final_items if len(r) > 1]
    return results if results else [text]

def normalize_text(text: str) -> tuple[str, int, str]:
    t = text.lower().strip()
    
    voice_corrections = [
        (r'\bpe\s+experien[tț][aăe]\b', 'pex-penta'),
        (r'\bpex\s*penta\b', 'pex-penta'),
        (r'\bcal\s+de\b', 'kalde'),
        (r'\b[sș]efu(\s*[\-]?\s*[iî]n)?\b', 'sifon'),
        (r'\b[sș]ef\s+in\b', 'sifon'),
        (r'\bval\s*rom\b', 'valrom'),
        (r'\b[tț]evi\b', 'teava'),
        (r'\bramifica[tț]ii\b', 'ramificatie'),
        (r'\baparent[aăe]?\b', 'aparent'),
        (r'\bsp[aă]lat\b', 'spalat'),
        (r'\bpardoseal[aă]\b', 'pardoseala'),
        (r'\btreisferturi\b', '3/4'),
        (r'\btrei\s*sferturi\b', '3/4'),
        (r'\b(un|unu)\s*sfert\b', '1/4'),
    ]
    for pattern, repl in voice_corrections:
        t = re.sub(pattern, repl, t, flags=re.IGNORECASE)

    cantitate = 1
    um = "buc"
    
    qty_match = re.search(r'\b(\d+)\s*(buc(?:ati|ăți)?|m(?:etri)?|ml|set(?:uri)?|rol(?:e|a)?|cut(?:ii|ie)?)\b', t)
    if qty_match:
        cantitate = int(qty_match.group(1))
        unit = qty_match.group(2)
        if "m" in unit:
            um = "m"
        elif "set" in unit:
            um = "set"
        elif "rol" in unit:
            um = "rola"
        elif "cut" in unit:
            um = "cutie"
        else:
            um = "buc"
        t = t.replace(qty_match.group(0), " ")
    else:
        simple_qty = re.match(r'^(\d+)\s+(.+)$', t)
        if simple_qty and int(simple_qty.group(1)) <= 2000:
            cantitate = int(simple_qty.group(1))
            t = simple_qty.group(2)

    replacements = [
        (r'\b(de\s+la|firma|marca|producator|pentru|cu|de)\b', ' '),
        (r'\b(unu|un|o)\s+p\s+2\b', '1/2'),
        (r'\b(unu|un|o)\s+pe\s+doi\b', '1/2'),
        (r'\bo\s+doime\b', '1/2'),
        (r'\b1\s+pe\s+2\b', '1/2'),
        (r'\b1\s+p\s+2\b', '1/2'),
        (r'\bjum[aă]tate\b', '1/2'),
        
        (r'\btrei\s+p\s+4\b', '3/4'),
        (r'\btrei\s+pe\s+patru\b', '3/4'),
        (r'\btrei\s+patrimi\b', '3/4'),
        (r'\b3\s+pe\s+4\b', '3/4'),
        (r'\b3\s+p\s+4\b', '3/4'),

        (r'\btrei\s+p\s+8\b', '3/8'),
        (r'\btrei\s+pe\s+opt\b', '3/8'),
        (r'\btrei\s+optimi\b', '3/8'),
        (r'\b3\s+pe\s+8\b', '3/8'),
        (r'\b3\s+p\s+8\b', '3/8'),

        (r'\b(unu|un)\s+p\s+4\b', '1/4'),
        (r'\b(unu|un)\s+pe\s+patru\b', '1/4'),
        (r'\bo\s+patrime\b', '1/4'),
        (r'\b1\s+pe\s+4\b', '1/4'),
        (r'\b1\s+p\s+4\b', '1/4'),

        (r'\b(un|unu)\s+(tol|tzol|zol)\s+si\s+un\s+sfert\b', '1 1/4'),
        (r'\b(un|unu)\s+(tol|tzol|zol)\s+si\s+jum[aă]tate\b', '1 1/2'),
        (r'\b(un|unu)\s+(tol|tzol|zol)\b', '1*'),
        (r'\bdoi\s+(toli|tzoli|zoli)\b', '2*'),
        (r'\btrei\s+(toli|tzoli|zoli)\b', '3*'),
        (r'\binel(?:e)?\s+alunec(?:ator|atoare)?\b', 'manson tece'),
        (r'\b[ϕφ]\s*(\d+)\b', r'\1'),
        (r'\b(?:fi|d)\s*(\d+)\b', r'\1'),
        (r'\b(gondor|glandez)\b', 'olandez'),
        (r'\bcaseta\s+1200\b', 'caseta metal 1200 distribuitor'),
    ]

    for pattern, repl in replacements:
        t = re.sub(pattern, repl, t, flags=re.IGNORECASE)

    t = re.sub(r'\s+', ' ', t).strip()
    return t, cantitate, um

LEARNED_FILE = os.path.join(os.path.dirname(__file__), "learned_mappings.json")

def load_learned_mappings() -> dict:
    if os.path.exists(LEARNED_FILE):
        try:
            with open(LEARNED_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print(f"Eroare citire memorie: {e}")
            return {}
    return {}

def clean_learn_key(text: str) -> str:
    t = text.lower()
    t = re.sub(r'^\[(?:foto|dictat)\]\s*', '', t)
    t = re.sub(r'^\d+[\s\.\-]+(?:buc|m|cut|set|colac|cutie|role)?\s*', '', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def save_learned_mapping(raw_query: str, product_data: dict) -> bool:
    key = clean_learn_key(raw_query)
    if not key or len(key) < 3:
        return False
    
    mappings = load_learned_mappings()
    mappings[key] = {
        "denumire": product_data.get("denumire"),
        "cod": product_data.get("cod", "-"),
        "cod_extern": product_data.get("cod_extern", "-"),
        "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }
    
    try:
        with open(LEARNED_FILE, "w", encoding="utf-8") as f:
            json.dump(mappings, f, ensure_ascii=False, indent=2)
        print(f"-> [MEMORIE] Asociere retinuta: '{key}' -> {product_data.get('denumire')}")
        return True
    except Exception as e:
        print(f"Eroare salvare memorie: {e}")
        return False

def check_learned_mapping(clean_query: str) -> dict | None:
    mappings = load_learned_mappings()
    if not mappings:
        return None
        
    q_key = clean_learn_key(clean_query)
    if not q_key:
        return None
        
    # 1. Potrivire exactă
    if q_key in mappings:
        return mappings[q_key]
        
    # 2. Potrivire fuzzy strânsă (>= 90%)
    best_item = None
    best_score = 0
    for k, v in mappings.items():
        s = fuzz.token_set_ratio(q_key, k)
        if s > best_score:
            best_score = s
            best_item = v
            
    if best_score >= 90:
        return best_item
        
    return None

def search_product(dictated_text: str, allow_fallback: bool = False) -> dict:
    clean_query, qty, um = normalize_text(dictated_text)
    
    # 0. Verificare mai întâi în memoria de învățare a aplicației!
    learned = check_learned_mapping(clean_query) or check_learned_mapping(dictated_text)
    if learned:
        return {
            "found": True,
            "is_ambiguous": False,
            "was_learned": True,
            "query_original": dictated_text,
            "query_normalized": clean_query,
            "cantitate": qty,
            "um": um,
            "best_match": {
                "denumire": learned["denumire"],
                "cod": learned.get("cod", "-"),
                "cod_extern": learned.get("cod_extern", "-"),
                "score": 100
            },
            "other_candidates": []
        }
    
    clean_query_spaced = re.sub(r'\b(\d+)[\s\-]+(\d+)\b', r'\1 \2', clean_query)
    tech_numbers = set(re.findall(r'(?:\d+\/\d+|\d+\*|\b\d+\b)', clean_query_spaced))
    
    # Eliminăm cantitatea extrasă dacă a fost găsită printre numere
    if str(qty) in tech_numbers and len(tech_numbers) > 1:
        tech_numbers.remove(str(qty))
        
    words = re.findall(r'[a-zA-Zăîșțâ]{3,}', clean_query)
    stop_words = {'din', 'de', 'la', 'cu', 'pentru', 'grade', 'grad', 'gr', 'si', 'pe'}
    key_words = [normalize_plurals(w) for w in words if w not in stop_words and w not in ('pvc', 'pp')]
    
    candidates = []
    
    for item in NOMENCLATOR:
        den = item["denumire"].upper()
        den_nums = set(re.findall(r'(?:\d+\/\d+|\d+\*|\b\d+\b)', den))
        
        # Filtru 1: Numerele tehnice cerute de utilizator (ex: 40 și 45 sau 1/2) TREBUIE să fie prezente
        missing_num = False
        for n in tech_numbers:
            pat = r'(?<!\d)' + re.escape(n) + r'(?!\d)'
            if not re.search(pat, den):
                missing_num = True
                break
        if missing_num and len(tech_numbers) > 0:
            continue
            
        # Filtru 2: Cuvinte cheie de produs
        matched_kw = 0
        for kw in key_words:
            root = kw[:min(len(kw), 5)]
            if root in den:
                matched_kw += 1
        
        if key_words and matched_kw == 0:
            continue
            
        score = fuzz.token_set_ratio(clean_query.upper(), den)
        
        # Bonus dacă conține toate cuvintele cheie
        if key_words and matched_kw == len(key_words):
            score += 25

        # Brand bonus / penalizare
        major_brands = ['TECE', 'KALDE', 'PURMO', 'VALROM', 'FERRO', 'TIEMME', 'IVAR', 'EVER', 'GF']
        q_upper = clean_query.upper()
        for b in major_brands:
            if b in q_upper:
                if b in den:
                    score += 45
                else:
                    for other_b in major_brands:
                        if other_b != b and other_b in den:
                            score -= 30
            
        # Penalizare severă pentru diametre gigantice necerute (110, 160, 200) dacă utilizatorul nu a cerut 110/160
        extra_nums = den_nums - tech_numbers
        for n in extra_nums:
            if n in {'110', '160', '200', '125', '75', '32'}:
                score -= 40

        # Penalizare pentru racorduri și reducții dacă s-a cerut un simplu fiting
        if 'RACORD' in den and 'RACORD' not in clean_query.upper():
            score -= 30
        if ('REDUCTIE' in den or 'REDUS' in den) and ('REDUCTIE' not in clean_query.upper() and 'REDUS' not in clean_query.upper()):
            score -= 25

        # Penalizare pentru atribute speciale necerute
        for attr in ['DUBLA', 'REGLABILA', 'REGL', 'FONO']:
            if attr in den and attr not in clean_query.upper():
                score -= 30
            
        candidates.append({
            "denumire": item["denumire"],
            "cod": item["cod"],
            "cod_extern": item["cod_extern"],
            "score": score
        })

    # Fallback dacă nu a găsit cu filtre stricte
    if not candidates:
        for item in NOMENCLATOR:
            den = item["denumire"].upper()
            matched = sum(1 for kw in key_words if kw[:min(len(kw), 5)] in den)
            if matched > 0:
                score = fuzz.token_set_ratio(clean_query.upper(), den)
                if score >= 55:
                    candidates.append({
                        "denumire": item["denumire"],
                        "cod": item["cod"],
                        "cod_extern": item["cod_extern"],
                        "score": score
                    })

    candidates.sort(key=lambda x: x["score"], reverse=True)
    
    seen = set()
    unique = []
    for c in candidates:
        if c["denumire"] not in seen:
            seen.add(c["denumire"])
            unique.append(c)
        if len(unique) >= 60:
            break

    if not unique:
        if allow_fallback:
            broad_matches = []
            clean_upper = clean_query.upper()
            for item in NOMENCLATOR:
                den = item["denumire"].upper()
                s = fuzz.token_set_ratio(clean_upper, den)
                if s >= 35:
                    broad_matches.append({
                        "denumire": item["denumire"],
                        "cod": item["cod"],
                        "cod_extern": item["cod_extern"],
                        "score": s
                    })
            broad_matches.sort(key=lambda x: x["score"], reverse=True)
            seen_b = set()
            broad_unique = []
            for c in broad_matches:
                if c["denumire"] not in seen_b:
                    seen_b.add(c["denumire"])
                    broad_unique.append(c)
                if len(broad_unique) >= 30:
                    break
                    
            return {
                "found": True,
                "is_unmatched": True,
                "is_ambiguous": True,
                "query_original": dictated_text,
                "query_normalized": clean_query,
                "cantitate": qty,
                "um": um,
                "best_match": {
                    "denumire": f"[DE VERIFICAT] {clean_query.upper()}",
                    "cod": "-",
                    "cod_extern": "-",
                    "score": 0
                },
                "other_candidates": broad_unique
            }

        return {
            "found": False,
            "query_original": dictated_text,
            "query_normalized": clean_query,
            "cantitate": qty,
            "um": um,
            "message": "Niciun reper găsit în nomenclator."
        }

    best = unique[0]
    second_score = unique[1]["score"] if len(unique) > 1 else 0

    is_ambiguous = False
    if len(unique) > 1:
        # Dacă există mai multe variante echivalente (scor apropiat) sau nu s-a specificat marca/materialul
        if (best["score"] - second_score < 12) or (best["score"] < 80):
            is_ambiguous = True

    return {
        "found": True,
        "query_original": dictated_text,
        "query_normalized": clean_query,
        "cantitate": qty,
        "um": um,
        "is_ambiguous": is_ambiguous,
        "best_match": best,
        "other_candidates": unique[1:]
    }

def parse_with_gemini(text: str) -> list[dict] | None:
    if not GEMINI_CLIENT:
        return None
    prompt = f"""Ești asistentul tehnic expert pentru depozitul de instalații sanitare Interconti.
Utilizatorul dictează prin voce pe telefon repere dorite pentru o comandă:
"{text}"

Ghid de interpretare depozit sanitar:
- "PVC" de interior (gri, subțire, scurgere) se referă în nomenclatorul depozitului la piese de scurgere din "PP" (polipropilenă) sau PVC.
- Păstrează strict dimensiunile (ex: 40, 50, 110) și unghiurile (ex: 45, 67, 87, 90). Nu adăuga 'dublă', 'reglabilă' sau alte diametre dacă nu s-au cerut.
- "pex-penta" / "pex penta" / "purmo" = TEAVA PEX-PENTA PURMO.
- "kalde" = KALDE (fitinguri alamă sau PPR).
- "jumătate" = 1/2, "trei sferturi" = 3/4, "un tol" = 1*.
- Recunoaște cantitățile (ex: 10 bucăți, 2 colaci, 120 metri) și unitatea de măsură (buc, m, colac, set, cutie). Dacă nu se precizează, cantitatea e 1 buc.

Pentru fiecare produs dictat, extrage:
- "termen_nomenclator": expresia tehnică de căutare optimă pentru denumirea din catalog (ex: "RAMIFICATIE PP 40 45", "TEAVA PEX-PENTA PURMO 16", "KALDE NIPLU 1/2")
- "cantitate": număr întreg sau zecimal
- "um": unitatea de măsură ("buc", "m", "colac", "set", etc.)

Răspunde STRICT sub formă de listă JSON validă:
[
  {{
    "termen_nomenclator": "RAMIFICATIE PP 40 45",
    "cantitate": 1,
    "um": "buc"
  }}
]
"""
    try:
        from google.genai import types
        import time
        models_to_try = ['gemini-3.6-flash', 'gemini-3.7-flash', 'gemini-3.5-flash-lite']
        data = None
        for m_name in models_to_try:
            try:
                resp = GEMINI_CLIENT.models.generate_content(
                    model=m_name,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.1
                    )
                )
                clean_text = resp.text.strip()
                if clean_text.startswith("```json"):
                    clean_text = clean_text[7:]
                elif clean_text.startswith("```"):
                    clean_text = clean_text[3:]
                if clean_text.endswith("```"):
                    clean_text = clean_text[:-3]
                clean_text = clean_text.strip()
                parsed = json.loads(clean_text)
                if isinstance(parsed, list) and len(parsed) > 0:
                    data = parsed
                    break
            except Exception as e_m:
                time.sleep(0.5)
        if data:
            return data
    except Exception as e:
        print(f"-> [AI] Eroare la procesarea cu Gemini: {e}")
    return None

def search_dictated_speech(text: str) -> list[dict]:
    # 1. Încercare cu Gemini AI (înțelegere contextuală, corecție fonetică, limbaj colocvial de șantier)
    gemini_items = parse_with_gemini(text)
    if gemini_items:
        results = []
        for g_item in gemini_items:
            q_term = g_item.get("termen_nomenclator", "")
            qty = g_item.get("cantitate", 1)
            um = g_item.get("um", "buc")
            res = search_product(q_term, allow_fallback=True)
            res["cantitate"] = qty
            res["um"] = um
            res["query_original"] = q_term or text
            if res.get("is_unmatched"):
                res["best_match"]["denumire"] = f"[DE VERIFICAT] {(q_term or text).upper()}"
            results.append(res)
        if results:
            print(f"-> [AI] Gemini a identificat cu succes {len(results)} repere din textul dictat!")
            return results

    # 2. Fallback de rezervă: algoritmul local fără internet (regex + fuzzy search)
    print("-> [Local] Folosire algoritm local de căutare...")
    segments = parse_multi_item_dictation(text)
    results = []
    for seg in segments:
        res = search_product(seg, allow_fallback=True)
        results.append(res)
    return results


def parse_image_with_gemini(image_bytes: bytes, mime_type: str = "image/jpeg") -> list[dict] | None:
    if not GEMINI_CLIENT:
        return None
    from google.genai import types
    import time
    
    prompt = """Ești expertul tehnic în instalații sanitare și termice pentru depozitul Interconti (Suceava).
În imagine este o comandă internă / listă de șantier scrisă de mână pe formular tipizat Interconti sau pe foaie de caiet.

REGULI CRUCIALE DE CITIRE PENTRU FORMULARELE DE INSTALAȚII:
1. Simbolul cerc tăiat 'ϕ' sau 'φ' (pe care OCR-ul îl confundă uneori cu 'p') urmat de numere înseamnă DIAMETRU / FI (ex: 'ϕ 16' = 16, 'ϕ 20' = 20, 'ϕ 26' = 26). Nu scrie niciodată 'p 16' sau 'p 20'! Scrie direct '16' sau '20' sau '26'!
2. Brandul principal de fitinguri și țevi cu manșon prin alunecare din listă este 'TECE' (scris adesea 'Tece' sau 'Tece.').
3. 'Inel alunecator' sau 'Inel alunec.' în sistemul Teceflex este MANSON TECE (ex: MANSON TECE 16, MANSON TECE 20, MANSON TECE 25).
4. 'Teuri' (scris cu T mare caligrafic) = TECE TEU ALAMA (ex: TECE TEU ALAMA 16).
5. 'Coturi ... F Tece' = TECE COT CU TALPA ALAMA 16-1/2 FI sau TECE COT ALAMA 16-1/2 FI.
6. 'Cot ... Tece' = TECE COT ALAMA (ex: TECE COT ALAMA 20-3/4 FE, TECE COT ALAMA 20-1/2 FI, TECE COT ALAMA 25-3/4 FI, TECE COT ALAMA 25-1 FE).
7. 'Teava fi 16 Tece rosie' = TECE TEAVA MULTISTRAT COPEX ROSU 16 (unitate de măsură: m).
8. 'Teava fi 16 Tece albastra' = TECE TEAVA MULTISTRAT COPEX ALBASTRU 16 (unitate de măsură: m).
9. 'Robinet cu Olandez 1"' / 'Robinet cu Olandez 3/4"' (litera 'Ol' caligrafică nu este 'Gondor', ci Olandez!).
10. 'Dop proba' / 'dop 1/2 proba' = CAL DOP PROBA 1/2 (dopuri de plastic pentru probă instalație).
11. 'Caseta 1200' = CASETA METAL 1200 DISTRIBUITOR (sau dulap distribuitor).
12. Dacă pe un rând sunt mai multe repere (ex: 'Cui beton = 1 cut', 'Disc = 1 buc', 'Caseta 1200 = 1'), separă-le obligatoriu ca articole distincte!
13. Include obligatoriu informațiile din coloana 'Observații' în termenul tehnic (ex: 'Distribuitor modular 3/4 5 cai rece', 'Distribuitor modular 3/4 3 cai cald', 'Distribuitor tur-retur Purmo 9 circuite').
14. Filet: 'M' = FE (Filet Exterior / tată), 'F' = FI (Filet Interior / mamă), 'MF' = mamă-tată, 'FF' = mamă-mamă.

Pentru fiecare produs identificat pe foaie, extrage:
- "termen_nomenclator": expresia tehnică optimă pentru căutare în catalog
- "cantitate": număr întreg sau zecimal
- "um": unitatea de măsură ("buc", "m", "colac", "cut", "set", etc.)
- "text_extras": textul exact descifrat de pe foaie (ex: "5 coturi 16 x 1/2 F Tece")

Răspunde STRICT sub formă de listă JSON validă:
[
  {
    "termen_nomenclator": "CUI BETON",
    "cantitate": 1,
    "um": "cut",
    "text_extras": "Cui beton = 1 cut"
  }
]
"""
    try:
        part = types.Part.from_bytes(data=image_bytes, mime_type=mime_type)
        
        # Ordine optimă: gemini-3.5-flash-lite este cel mai rapid și stabil pentru OCR, urmat de gemini-3.6-flash și 3.7
        models_to_try = ['gemini-3.5-flash-lite', 'gemini-3.6-flash', 'gemini-3.7-flash']
        data = None
        for m_name in models_to_try:
            for attempt in range(2):
                try:
                    resp = GEMINI_CLIENT.models.generate_content(
                        model=m_name,
                        contents=[part, prompt],
                        config=types.GenerateContentConfig(
                            response_mime_type="application/json",
                            temperature=0.1
                        )
                    )
                    clean_text = resp.text.strip()
                    if clean_text.startswith("```json"):
                        clean_text = clean_text[7:]
                    elif clean_text.startswith("```"):
                        clean_text = clean_text[3:]
                    if clean_text.endswith("```"):
                        clean_text = clean_text[:-3]
                    clean_text = clean_text.strip()

                    parsed = json.loads(clean_text)
                    if isinstance(parsed, list) and len(parsed) > 0:
                        data = parsed
                        print(f"-> [AI Vision] Succes descifrare imagine cu modelul {m_name} ({len(data)} repere gasite pe foaie)!")
                        break
                except Exception as e_m:
                    print(f"-> [AI Vision] Modelul {m_name} (incercarea {attempt+1}) eroare: {e_m}")
                    time.sleep(1.5)
            if data:
                break

        if data:
            return data
    except Exception as e:
        print(f"-> [AI Vision] Eroare generala la procesarea imaginii: {e}")
    return None


def search_image_order(image_bytes: bytes, mime_type: str = "image/jpeg") -> list[dict]:
    gemini_items = parse_image_with_gemini(image_bytes, mime_type)
    if not gemini_items:
        return []
    
    results = []
    for g_item in gemini_items:
        q_term = g_item.get("termen_nomenclator", "")
        qty = g_item.get("cantitate", 1)
        um = g_item.get("um", "buc")
        original_text = g_item.get("text_extras", q_term)
        
        # Înlocuim simbolurile grecești de diametru (phi) cu 'fi' pentru compatibilitate totală
        original_text = original_text.replace('φ', 'fi').replace('ϕ', 'fi').replace('Φ', 'fi')
        q_term = q_term.replace('φ', 'fi').replace('ϕ', 'fi').replace('Φ', 'fi')
        
        # Căutăm cu fallback activat obligatoriu!
        res = search_product(q_term, allow_fallback=True)
        res["cantitate"] = qty
        res["um"] = um
        res["query_original"] = f"[Foto] {original_text}"
        
        # Dacă reperul nu s-a regăsit în catalog, păstrăm numele original citit din poză
        if res.get("is_unmatched"):
            res["best_match"]["denumire"] = f"[DE VERIFICAT] {original_text.upper()}"
            
        results.append(res)

    print(f"-> [AI Vision] Returnate {len(results)} repere (inclusiv cele pentru verificare manuala)!")
    return results


def generate_word_document(items: list, output_path: str = None) -> str:
    comenzi_dir = os.path.join(os.path.dirname(__file__), "comenzi")
    if not output_path:
        os.makedirs(comenzi_dir, exist_ok=True)
        now_str = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        output_path = os.path.join(comenzi_dir, f"Comanda_{now_str}.docx")

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("INTERCONTI — NOTĂ DE COMANDĂ DEPOZIT")
    r_title.bold = True
    r_title.font.size = Pt(15)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(10)
    now = datetime.now()
    r_info = p_meta.add_run(f"Data: {now.strftime('%d.%m.%Y')}  |  Ora: {now.strftime('%H:%M')}  |  Total articole: {len(items)}\n")
    r_info.font.size = Pt(9.5)
    r_sub = p_meta.add_run("Notă: Articolele marcate cu [ ! ] necesită o scurtă privire de confirmare.")
    r_sub.font.size = Pt(8.5)
    r_sub.font.italic = True

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Inches(0.4), Inches(1.3), Inches(3.4), Inches(0.6), Inches(0.5), Inches(1.3)]
    headers = ["Nr.", "Cod Articol", "Denumire Reper Nomenclator", "Cant.", "U.M.", "Verificare / Obs"]

    hdr_cells = table.rows[0].cells
    for idx, (title, width) in enumerate(zip(headers, widths)):
        cell = hdr_cells[idx]
        cell.width = width
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(9)
        shading = parse_xml(r'<w:shd {} w:fill="222222"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for i, it in enumerate(items, 1):
        row_cells = table.add_row().cells
        for idx, width in enumerate(widths):
            row_cells[idx].width = width

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run(str(i)).font.size = Pt(9)

        cod_val = it.get("cod_extern") or it.get("cod") or "-"
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run(cod_val).font.size = Pt(8.5)

        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p2.add_run(it.get("denumire", ""))
        r2.bold = True
        r2.font.size = Pt(9)

        p3 = row_cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(str(it.get("cantitate", 1)))
        r3.bold = True
        r3.font.size = Pt(9.5)

        p4 = row_cells[4].paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p4.add_run(it.get("um", "buc")).font.size = Pt(8.5)

        p5 = row_cells[5].paragraphs[0]
        p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if it.get("is_ambiguous"):
            r5 = p5.add_run("[ ! ] VERIFICĂ")
            r5.bold = True
            r5.font.size = Pt(9)
        else:
            r5 = p5.add_run("[   ] OK")
            r5.font.size = Pt(8.5)

        for cell in row_cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'''
                <w:tcBorders {} >
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                    <w:top w:val="none"/>
                    <w:left w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            '''.format(nsdecls('w')))
            tcPr.append(tcBorders)

    p_footer = doc.add_paragraph()
    p_footer.paragraph_format.space_before = Pt(30)
    p_footer.add_run("Semnătură întocmit: ______________________             Semnătură eliberat depozit: ______________________").font.size = Pt(9)

    doc.save(output_path)
    return output_path

def generate_whatsapp_link(items: list, doc_download_url: str = "") -> str:
    now = datetime.now()
    lines = [
        f"📦 *COMANDĂ NOUĂ INTERCONTI* ({now.strftime('%d.%m %H:%M')})",
        f"Salut Bogdan, mai jos ai reperele pentru tipar:\n"
    ]
    for i, it in enumerate(items, 1):
        flag = " ⚠️ [VERIFICĂ]" if it.get("is_ambiguous") else ""
        cod = f" (Cod: {it.get('cod_extern') or it.get('cod')})" if (it.get('cod_extern') or it.get('cod')) else ""
        lines.append(f"{i}. {it.get('denumire')}{cod} — *{it.get('cantitate')} {it.get('um')}*{flag}")
    
    if doc_download_url:
        lines.append(f"\n📄 *Descarcă Documentul Word (.docx) gata de printat:*")
        lines.append(doc_download_url)
    
    msg = "\n".join(lines)
    encoded_msg = urllib.parse.quote(msg)
    return f"https://wa.me/{BOGDAN_PHONE}?text={encoded_msg}"
