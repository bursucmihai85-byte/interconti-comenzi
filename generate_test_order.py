import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

items = [
    {
        "cod": "5902650652132",
        "denumire": "ARMAKAN RAMIFICATIE PP 50-50-45",
        "cantitate": 1,
        "um": "buc",
        "obs": "Dictat: 'Ramificații pvc de 50-45 °'"
    },
    {
        "cod": "77161630",
        "denumire": "TECE TEAVA PT PARDOSEALA 16,COLAC 300M",
        "cantitate": 120,
        "um": "m",
        "obs": "Dictat: 'pe experienta (PEX) 16 120 m'"
    },
    {
        "cod": "2300000001396",
        "denumire": "- KALDE NIPLU ALAMA 1/2",
        "cantitate": 3,
        "um": "buc",
        "obs": "Dictat: 'niplu 1/2 cal de (Kalde) 3 buc'"
    },
    {
        "cod": "5947041008839",
        "denumire": "EUR SIFON APARENT DUBLU PT MASINA SPALAT+USCATOR IESIRE32",
        "cantitate": 2,
        "um": "buc",
        "obs": "Dictat: 'șefu în (sifon) masina spalat aparent 2 buc'"
    }
]

doc_path = r"C:\Users\Lucru\Desktop\interconti\Comanda_Test_Interconti.docx"
doc = Document()

# Setare margini (1.5 cm)
for section in doc.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# Antet
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_title.add_run("INTERCONTI — NOTĂ DE COMANDĂ DEPOZIT")
r_title.bold = True
r_title.font.size = Pt(16)
r_title.font.color.rgb = RGBColor(16, 44, 87) # Albastru marin

p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_after = Pt(12)
r_date = p_meta.add_run(f"Data & Ora emiterii: {datetime.now().strftime('%d.%m.%Y - %H:%M')}\n")
r_date.font.size = Pt(9.5)
r_items_count = p_meta.add_run(f"Total repere pe comandă: {len(items)} articole  |  Mod: Preluare Automată prin Dictare")
r_items_count.bold = True
r_items_count.font.size = Pt(9.5)

# Tabel
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

widths = [Inches(0.4), Inches(1.3), Inches(3.2), Inches(0.6), Inches(0.5), Inches(1.6)]
headers = ["Nr.", "Cod Articol", "Denumire Nomenclator", "Cant.", "U.M.", "Observații / Sursă Dictare"]

# Stilizare header tabel
hdr_cells = table.rows[0].cells
for idx, (title, width) in enumerate(zip(headers, widths)):
    cell = hdr_cells[idx]
    cell.width = width
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(255, 255, 255)
    shading = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)

# Rânduri tabel
for i, it in enumerate(items, 1):
    row_cells = table.add_row().cells
    for idx, width in enumerate(widths):
        row_cells[idx].width = width

    # Nr.
    p0 = row_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.add_run(str(i)).font.size = Pt(9)

    # Cod
    p1 = row_cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run(it["cod"]).font.size = Pt(8.5)

    # Denumire
    p2 = row_cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run(it["denumire"])
    r2.bold = True
    r2.font.size = Pt(9)

    # Cantitate
    p3 = row_cells[3].paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(str(it["cantitate"]))
    r3.bold = True
    r3.font.size = Pt(9.5)

    # U.M.
    p4 = row_cells[4].paragraphs[0]
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.add_run(it["um"]).font.size = Pt(8.5)

    # Observații
    p5 = row_cells[5].paragraphs[0]
    p5.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r5 = p5.add_run(it["obs"])
    r5.font.size = Pt(8)
    r5.font.italic = True

    # Zebra striping
    if i % 2 == 0:
        for cell in row_cells:
            shd = parse_xml(r'<w:shd {} w:fill="F3F4F6"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shd)

# Subsol
p_space = doc.add_paragraph()
p_space.paragraph_format.space_before = Pt(25)

p_sign = doc.add_paragraph()
p_sign.add_run("Întocmit de: ___________________________            Predat / Pregătit de: ___________________________").font.size = Pt(9.5)

doc.save(doc_path)
print(f"Document generat cu succes la: {doc_path}")
